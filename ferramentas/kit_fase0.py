# -*- coding: utf-8 -*-
"""Gera o Kit Fase 0 — Diagnóstico (xlsx) do projeto de automação fiscal."""
import datetime as _dt
import json
import sys

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

CATALOGO_JSON = sys.argv[1]
SAIDA = sys.argv[2]

FONTE = "Arial"
VERDE = "1D6B4F"
VERDE_CLARO = "E7F0E3"
AMARELO = "FFF2CC"
AMARELO_FORTE = "FFFF00"
CINZA = "BFC9C2"

f_titulo = Font(name=FONTE, size=15, bold=True, color=VERDE)
f_sub = Font(name=FONTE, size=10, italic=True, color="55635B")
f_secao = Font(name=FONTE, size=12, bold=True, color=VERDE)
f_head = Font(name=FONTE, size=10, bold=True, color="FFFFFF")
f_norm = Font(name=FONTE, size=10)
f_bold = Font(name=FONTE, size=10, bold=True)
f_input = Font(name=FONTE, size=10, color="0000FF")
f_link = Font(name=FONTE, size=10, color="008000")
f_ex = Font(name=FONTE, size=10, italic=True, color="7F7F7F")

fill_head = PatternFill("solid", fgColor=VERDE)
fill_input = PatternFill("solid", fgColor=AMARELO)
fill_key = PatternFill("solid", fgColor=AMARELO_FORTE)
fill_band = PatternFill("solid", fgColor=VERDE_CLARO)

borda = Border(*(Side(style="thin", color=CINZA),) * 4)
wrap = Alignment(wrap_text=True, vertical="top")
centro = Alignment(horizontal="center", vertical="center", wrap_text=True)

BRL = 'R$ #,##0;[Red]-R$ #,##0;"-"'
PCT = '0.0%;-0.0%;"-"'
NUM1 = '#,##0.0;-#,##0.0;"-"'
NUM0 = '#,##0;-#,##0;"-"'

ETAPAS = [
    ("1. Coleta e cobrança de documentos",
     "Baixar XML e relatórios, cobrar o cliente, organizar pastas, receber documentos."),
    ("2. Importação e escrituração",
     "Importar XML/TXT na Domínio, lançar notas manualmente, configurar acumuladores e produtos, corrigir erros de importação."),
    ("3. Conferência e auditoria",
     "Conferir notas × escrita, CFOP/CST, totais, relatórios de conferência e cruzamentos."),
    ("4. Apuração e guias",
     "Calcular tributos na Domínio, recalcular, emitir guias (DAS, DARF, estaduais, ISS)."),
    ("5. Obrigações acessórias",
     "Gerar, validar e transmitir SPEDs, PGDAS-D, DCTFWeb/MIT, EFD-Reinf e declarações estaduais/municipais."),
    ("6. Atendimento e envio ao cliente",
     "Enviar guias e relatórios, responder dúvidas, reuniões com clientes."),
    ("7. Outros",
     "Reuniões internas, treinamento, suporte de TI, tarefas não fiscais."),
]
PCT_TEMPO = [0.14, 0.32, 0.11, 0.12, 0.12, 0.11, 0.08]
POTENCIAL = {  # potencial de redução de tempo em 12 meses, por cenário
    "Conservador": [0.35, 0.18, 0.30, 0.15, 0.10, 0.20, 0.0],
    "Base":        [0.50, 0.25, 0.45, 0.25, 0.15, 0.30, 0.0],
    "Otimista":    [0.65, 0.35, 0.55, 0.35, 0.20, 0.40, 0.0],
}

UFS = "AC,AL,AM,AP,BA,CE,DF,ES,GO,MA,MG,MS,MT,PA,PB,PE,PI,PR,RJ,RN,RO,RR,RS,SC,SE,SP,TO"
N_EMPRESAS = 500   # linhas preparadas no cadastro
N_TEMPO = 2000     # linhas preparadas no levantamento
N_EQUIPE = 20

wb = Workbook()


def larguras(ws, mapa):
    for col, w in mapa.items():
        ws.column_dimensions[col].width = w


def cabecalho(ws, linha, textos, altura=30):
    for i, t in enumerate(textos, start=1):
        c = ws.cell(row=linha, column=i, value=t)
        c.font, c.fill, c.alignment, c.border = f_head, fill_head, centro, borda
    ws.row_dimensions[linha].height = altura


def titulo(ws, texto, sub=None):
    ws["A1"] = texto
    ws["A1"].font = f_titulo
    if sub:
        ws["A2"] = sub
        ws["A2"].font = f_sub


def lista(ws, formula, rng, prompt=None):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    if prompt:
        dv.promptTitle, dv.prompt, dv.showInputMessage = "Preenchimento", prompt, True
    dv.error, dv.errorTitle, dv.showErrorMessage = "Escolha um valor da lista.", "Valor inválido", True
    ws.add_data_validation(dv)
    dv.add(rng)


def pinta_input(ws, rng):
    for row in ws[rng]:
        for c in row:
            c.fill, c.border = fill_input, borda
            if c.font is None or c.font == Font():
                c.font = f_norm


# ======================================================================
# 1. Leia-me
# ======================================================================
ws = wb.active
ws.title = "Leia-me"
titulo(ws, "Kit Fase 0 — Diagnóstico do Departamento Fiscal",
       "Projeto: Ecossistema de Automação Fiscal (Domínio Web) · versão 1.0 · setembro/2026")
larguras(ws, {"A": 34, "B": 110})
linhas = [
    ("PARA QUE SERVE", None),
    ("Objetivo", "Medir o ponto de partida (baseline) e organizar os dados mestres antes de automatizar. "
                 "Sem esses números não dá para priorizar nem provar a redução de mão de obra."),
    ("Quando usar", "Nas 4 semanas da Fase 0. O levantamento de tempo dura 10 dias úteis; o cadastro mestre deve cobrir 100% das empresas."),
    ("ORDEM DE PREENCHIMENTO", None),
    ("Passo 1 — Equipe", "Liste os analistas do fiscal. Os nomes alimentam as listas das outras abas."),
    ("Passo 2 — Cadastro Mestre", "Uma linha por empresa (cada analista preenche a sua carteira). A classe A/B/C é sugerida automaticamente."),
    ("Passo 3 — Levantamento de Tempo", "Cada analista registra suas atividades por 10 dias úteis (blocos de 15 min ou mais), escolhendo uma das 7 etapas."),
    ("Passo 4 — Resumo Baseline", "Calcula sozinho: horas por etapa, % de retrabalho, horas por empresa e por classe. Informe os dias úteis medidos."),
    ("Passo 5 — Business Case", "Troque as premissas (amarelo, texto azul) pelos números reais e escolha o cenário. Mostra FTE liberado, custo e payback."),
    ("Passo 6 — Catálogo de Auditorias", "46 cruzamentos priorizados por fase. Use as colunas Status/Responsável para acompanhar a construção."),
    ("Passo 7 — Plano 30 dias e Roteiro", "Checklist da Fase 0 e perguntas para Thomson Reuters, SERPRO e fornecedores."),
    ("LEGENDA DE CORES", None),
    ("Fundo amarelo claro", "Célula para você preencher."),
    ("Texto azul + fundo amarelo", "Premissa editável do Business Case (estimativa inicial; substitua pelo número real)."),
    ("Texto preto", "Fórmula. Não sobrescreva."),
    ("Texto verde", "Vínculo com outra aba."),
    ("Linha em cinza itálico", "EXEMPLO de preenchimento. Substitua ou apague antes de usar os resultados."),
    ("AS 7 ETAPAS DO TRABALHO FISCAL", None),
]
linhas += [(e, d) for e, d in ETAPAS]
linhas += [
    ("CLASSES DE EMPRESA", None),
    ("A — complexa", "Lucro Real, indústria ou mais de 500 documentos/mês."),
    ("B — intermediária", "Lucro Presumido, sujeita a ST ou 100 a 500 documentos/mês."),
    ("C — simples", "Demais casos (Simples Nacional de baixo volume, serviços, MEI)."),
    ("CUIDADOS", None),
    ("LGPD e sigilo", "Esta planilha tem dados de clientes e, na aba Equipe, custos de pessoal. Guarde em pasta com acesso restrito aos sócios e à coordenação."),
    ("Honestidade dos registros", "O levantamento de tempo não é avaliação individual. Ele mostra onde o processo trava. Diga isso à equipe antes de começar."),
]
r = 4
for a, b in linhas:
    if b is None:
        ws.cell(row=r, column=1, value=a).font = f_secao
        r += 1
        continue
    ws.cell(row=r, column=1, value=a).font = f_bold
    c = ws.cell(row=r, column=2, value=b)
    c.font, c.alignment = f_norm, wrap
    ws.cell(row=r, column=1).alignment = wrap
    if a.startswith("Fundo amarelo"):
        ws.cell(row=r, column=1).fill = fill_input
    elif a.startswith("Texto azul"):
        ws.cell(row=r, column=1).fill = fill_key
        ws.cell(row=r, column=1).font = Font(name=FONTE, size=10, bold=True, color="0000FF")
    elif a.startswith("Texto verde"):
        ws.cell(row=r, column=1).font = Font(name=FONTE, size=10, bold=True, color="008000")
    elif a.startswith("Linha em cinza"):
        ws.cell(row=r, column=1).font = Font(name=FONTE, size=10, bold=True, italic=True, color="7F7F7F")
    r += 1
ws.sheet_view.showGridLines = False

# ======================================================================
# 2. Equipe
# ======================================================================
ws = wb.create_sheet("Equipe")
titulo(ws, "Equipe do departamento fiscal",
       "Os nomes desta aba alimentam as listas do Cadastro Mestre e do Levantamento de Tempo. Custo é opcional e confidencial.")
cabecalho(ws, 4, ["Nome", "Cargo", "Função atual (texto livre)", "Custo mensal total (R$)",
                  "Perfil técnico", "Interesse em automação?"])
larguras(ws, {"A": 28, "B": 18, "C": 40, "D": 20, "E": 26, "F": 18})
ex = ["Maria Exemplo (EXEMPLO)", "Analista", "Carteira Simples Nacional — comércio", 5200,
      "Excel avançado / Power Query", "S"]
for i, v in enumerate(ex, start=1):
    c = ws.cell(row=5, column=i, value=v)
    c.font, c.border = f_ex, borda
ws.cell(row=5, column=4).number_format = BRL
ult_eq = 5 + N_EQUIPE - 1
pinta_input(ws, f"A6:F{ult_eq}")
for rr in range(6, ult_eq + 1):
    ws.cell(row=rr, column=4).number_format = BRL
lista(ws, '"Assistente,Analista,Analista sênior,Coordenador,Outro"', f"B5:B{ult_eq}")
lista(ws, '"Básico,Excel avançado,Excel avançado / Power Query,SQL,Python,Outro"', f"E5:E{ult_eq}")
lista(ws, '"S,N"', f"F5:F{ult_eq}")
ws.cell(row=ult_eq + 2, column=3, value="Custo médio informado (R$)").font = f_bold
ws.cell(row=ult_eq + 2, column=4, value=f"=IFERROR(AVERAGE(D6:D{ult_eq}),0)").number_format = BRL
ws.freeze_panes = "A5"
EQUIPE_RNG = f"Equipe!$A$5:$A${ult_eq}"
EQUIPE_MEDIA = f"Equipe!$D${ult_eq + 2}"

# ======================================================================
# 3. Cadastro Mestre
# ======================================================================
ws = wb.create_sheet("Cadastro Mestre")
cols = [
    ("Cód. Domínio", 11), ("CNPJ", 20), ("Razão social", 34), ("Regime", 18), ("Atividade", 13),
    ("UF", 6), ("Município", 18), ("Contribuinte ICMS?", 12), ("Sujeita a ST?", 10),
    ("NF-e/NFC-e saída/mês", 12), ("NF-e entrada/mês", 12), ("NFS-e/mês (prest.+tomadas)", 13),
    ("CT-e/mês", 9), ("Documentos/mês (auto)", 12), ("Classe sugerida (auto)", 11),
    ("Classe — ajuste manual", 11), ("Classe final (auto)", 10), ("Analista responsável", 24),
    ("Certificado", 11), ("Validade do certificado", 13), ("Dias p/ vencer (auto)", 11),
    ("Procuração e-CAC ao escritório?", 13), ("Validade da procuração", 13),
    ("PGDAS-D/DEFIS", 10), ("DCTFWeb/MIT", 10), ("EFD-Reinf", 10), ("EFD ICMS/IPI", 10),
    ("EFD-Contrib.", 10), ("Declaração estadual (qual?)", 16), ("ISS/declaração municipal", 13),
    ("Benefício fiscal / regime especial", 22), ("Importa XML na Domínio hoje?", 12),
    ("Observações", 34),
]
titulo(ws, "Cadastro Mestre de Empresas",
       "Uma linha por empresa. Colunas (auto) são fórmulas. A classe final usa o ajuste manual, se houver.")
HL = 4
cabecalho(ws, HL, [c for c, _ in cols], altura=44)
for i, (_, w) in enumerate(cols, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w
L0 = HL + 1                      # linha de exemplo
L1 = L0 + 1                      # primeira linha real
LN = L0 + N_EMPRESAS - 1         # última linha preparada
exemplo = [101, "00.000.000/0001-00", "Comércio Exemplo Ltda (EXEMPLO)", "Simples Nacional", "Comércio",
           "SP", "Campinas", "S", "S", 180, 95, 4, 6, None, None, None, None,
           "Maria Exemplo (EXEMPLO)", "A1", "2027-03-15", None, "S", "2027-08-31",
           "S", "S", "N", "N", "N", "DeSTDA", "N", "", "Parcial", "Linha de exemplo: substitua ou apague."]
for rr in range(L0, LN + 1):
    for ci in range(1, len(cols) + 1):
        c = ws.cell(row=rr, column=ci)
        c.border = borda
        c.font = f_ex if rr == L0 else f_norm
        if rr > L0 and ci not in (14, 15, 17, 21):
            c.fill = fill_input
    # fórmulas
    ws.cell(row=rr, column=14, value=f'=IF(C{rr}="","",SUM(J{rr}:M{rr}))').number_format = NUM0
    ws.cell(row=rr, column=15, value=(
        f'=IF(C{rr}="","",IF(OR(D{rr}="Lucro Real",E{rr}="Indústria",N{rr}>500),"A",'
        f'IF(OR(D{rr}="Lucro Presumido",I{rr}="S",N{rr}>=100),"B","C")))'))
    ws.cell(row=rr, column=17, value=f'=IF(C{rr}="","",IF(P{rr}<>"",P{rr},O{rr}))')
    ws.cell(row=rr, column=21, value=f'=IF(T{rr}="","",T{rr}-TODAY())').number_format = NUM0
    ws.cell(row=rr, column=20).number_format = "DD/MM/YYYY"
    ws.cell(row=rr, column=23).number_format = "DD/MM/YYYY"
    for ci in (10, 11, 12, 13):
        ws.cell(row=rr, column=ci).number_format = NUM0
    for ci in (14, 15, 17, 21):
        ws.cell(row=rr, column=ci).alignment = Alignment(horizontal="center")
for ci, v in enumerate(exemplo, start=1):
    if v is None or ci in (14, 15, 17, 21):
        continue
    if ci in (20, 23):
        v = _dt.datetime.strptime(v, "%Y-%m-%d")
    ws.cell(row=L0, column=ci, value=v)
rng = lambda col: f"{col}{L0}:{col}{LN}"
lista(ws, '"Simples Nacional,SIMEI (MEI),Lucro Presumido,Lucro Real,Imune/Isenta"', rng("D"))
lista(ws, '"Comércio,Indústria,Serviços,Misto"', rng("E"))
lista(ws, f'"{UFS}"', rng("F"))
for col in ("H", "I", "V", "X", "Y", "Z", "AA", "AB", "AD"):
    lista(ws, '"S,N"', rng(col))
lista(ws, '"A,B,C"', rng("P"), "Deixe vazio para usar a classe sugerida.")
lista(ws, f"={EQUIPE_RNG}", rng("R"))
lista(ws, '"A1,A3,Não possui"', rng("S"))
lista(ws, '"S,N,Parcial"', rng("AF"))
ws.conditional_formatting.add(f"U{L0}:U{LN}", FormulaRule(
    formula=[f'AND(ISNUMBER(U{L0}),U{L0}<15)'], fill=PatternFill("solid", fgColor="F4C7C3"),
    font=Font(name=FONTE, color="9C0006", bold=True)))
ws.conditional_formatting.add(f"U{L0}:U{LN}", FormulaRule(
    formula=[f'AND(ISNUMBER(U{L0}),U{L0}>=15,U{L0}<30)'], fill=PatternFill("solid", fgColor="FCE8B2")))
ws.freeze_panes = ws.cell(row=L0, column=4)
ws.auto_filter.ref = f"A{HL}:{get_column_letter(len(cols))}{LN}"
CAD_COD = f"'Cadastro Mestre'!$A${L0}:$A${LN}"
CAD_NOME = f"'Cadastro Mestre'!$C${L0}:$C${LN}"
CAD_CLASSE = f"'Cadastro Mestre'!$Q${L0}:$Q${LN}"
CAD_RESP = f"'Cadastro Mestre'!$R${L0}:$R${LN}"

# ======================================================================
# 4. Levantamento de Tempo
# ======================================================================
ws = wb.create_sheet("Levantamento de Tempo")
titulo(ws, "Levantamento de Tempo (10 dias úteis)",
       "Registre cada bloco de trabalho (mín. 15 min). Empresa e classe aparecem sozinhas a partir do código.")
cabecalho(ws, 4, ["Data", "Analista", "Cód. empresa", "Empresa (auto)", "Classe (auto)", "Etapa",
                  "Minutos", "Retrabalho?", "Detalhe (o que foi feito)"])
larguras(ws, {"A": 12, "B": 24, "C": 11, "D": 32, "E": 9, "F": 34, "G": 10, "H": 11, "I": 48})
T0, TN = 5, 5 + N_TEMPO - 1
for rr in range(T0, TN + 1):
    for ci in range(1, 10):
        c = ws.cell(row=rr, column=ci)
        c.border = borda
        c.font = f_ex if rr == T0 else f_norm
        if rr > T0 and ci not in (4, 5):
            c.fill = fill_input
    ws.cell(row=rr, column=1).number_format = "DD/MM/YYYY"
    ws.cell(row=rr, column=4, value=(
        f'=IF(C{rr}="","",IFERROR(INDEX({CAD_NOME},MATCH(C{rr},{CAD_COD},0)),"código não encontrado"))'))
    ws.cell(row=rr, column=5, value=(
        f'=IF(C{rr}="","",IFERROR(INDEX({CAD_CLASSE},MATCH(C{rr},{CAD_COD},0)),""))'))
    ws.cell(row=rr, column=5).alignment = Alignment(horizontal="center")
ex = [_dt.datetime(2026, 10, 1), "Maria Exemplo (EXEMPLO)", 101, None, None, ETAPAS[1][0], 45, "S",
      "Reimportação de 12 XML com acumulador errado (EXEMPLO)"]
for ci, v in enumerate(ex, start=1):
    if v is not None:
        ws.cell(row=T0, column=ci, value=v)
lista(ws, f"={EQUIPE_RNG}", f"B{T0}:B{TN}")
lista(ws, '"' + ",".join(e for e, _ in ETAPAS) + '"', f"F{T0}:F{TN}")
lista(ws, '"S,N"', f"H{T0}:H{TN}", "S = refazendo algo que já tinha sido feito (erro, devolução, correção).")
dv = DataValidation(type="whole", operator="between", formula1="1", formula2="600", allow_blank=True)
dv.error, dv.errorTitle = "Informe minutos entre 1 e 600.", "Minutos"
ws.add_data_validation(dv)
dv.add(f"G{T0}:G{TN}")
ws.freeze_panes = "A5"
ws.auto_filter.ref = f"A4:I{TN}"
LT_ANAL = f"'Levantamento de Tempo'!$B${T0}:$B${TN}"
LT_CLASSE = f"'Levantamento de Tempo'!$E${T0}:$E${TN}"
LT_ETAPA = f"'Levantamento de Tempo'!$F${T0}:$F${TN}"
LT_MIN = f"'Levantamento de Tempo'!$G${T0}:$G${TN}"
LT_RET = f"'Levantamento de Tempo'!$H${T0}:$H${TN}"

# ======================================================================
# 5. Resumo Baseline
# ======================================================================
ws = wb.create_sheet("Resumo Baseline")
titulo(ws, "Resumo do Baseline",
       "Calculado a partir do Levantamento de Tempo e do Cadastro Mestre. Preencha só as células amarelas.")
larguras(ws, {"A": 38, "B": 16, "C": 16, "D": 14, "E": 16, "F": 18})
ws["A4"], ws["A5"], ws["A6"] = "Dias úteis medidos", "Dias úteis em um mês", "Fator de projeção mensal"
for a in ("A4", "A5", "A6"):
    ws[a].font = f_bold
ws["B4"], ws["B5"] = 10, 21
for a in ("B4", "B5"):
    ws[a].font, ws[a].fill, ws[a].border = f_input, fill_key, borda
ws["B6"] = "=IF(B4=0,0,B5/B4)"
ws["B6"].number_format = "0.00"
ws["C4"] = "Premissa: 10 dias úteis de medição (2 semanas)."
ws["C4"].font = f_sub

ws["A8"] = "Tempo por etapa"
ws["A8"].font = f_secao
cabecalho(ws, 9, ["Etapa", "Minutos no período", "Horas no período", "% do total", "% de retrabalho",
                  "Horas/mês projetadas"])
E0 = 10
for i, (etapa, _) in enumerate(ETAPAS):
    rr = E0 + i
    ws.cell(row=rr, column=1, value=etapa).font = f_norm
    ws.cell(row=rr, column=2, value=f"=SUMIFS({LT_MIN},{LT_ETAPA},A{rr})").number_format = NUM0
    ws.cell(row=rr, column=3, value=f"=B{rr}/60").number_format = NUM1
    ws.cell(row=rr, column=4, value=f"=IF($B${E0 + 7}=0,0,B{rr}/$B${E0 + 7})").number_format = PCT
    ws.cell(row=rr, column=5, value=(
        f'=IF(B{rr}=0,0,SUMIFS({LT_MIN},{LT_ETAPA},A{rr},{LT_RET},"S")/B{rr})')).number_format = PCT
    ws.cell(row=rr, column=6, value=f"=C{rr}*$B$6").number_format = NUM1
    for ci in range(1, 7):
        ws.cell(row=rr, column=ci).border = borda
ET = E0 + 7
ws.cell(row=ET, column=1, value="Total").font = f_bold
ws.cell(row=ET, column=2, value=f"=SUM(B{E0}:B{ET - 1})").number_format = NUM0
ws.cell(row=ET, column=3, value=f"=SUM(C{E0}:C{ET - 1})").number_format = NUM1
ws.cell(row=ET, column=4, value=f"=SUM(D{E0}:D{ET - 1})").number_format = PCT
ws.cell(row=ET, column=5, value=(
    f'=IF(B{ET}=0,0,SUMIFS({LT_MIN},{LT_RET},"S")/B{ET})')).number_format = PCT
ws.cell(row=ET, column=6, value=f"=SUM(F{E0}:F{ET - 1})").number_format = NUM1
for ci in range(1, 7):
    ws.cell(row=ET, column=ci).border = borda
    ws.cell(row=ET, column=ci).font = f_bold
RESUMO_PCT_ROW0 = E0  # usado pelo Business Case

ws.cell(row=ET + 2, column=1, value="Tempo por analista").font = f_secao
cabecalho(ws, ET + 3, ["Analista (da aba Equipe)", "Horas no período", "Empresas na carteira",
                       "Horas/empresa/mês", "", ""])
A0 = ET + 4
for i in range(N_EQUIPE):
    rr = A0 + i
    ws.cell(row=rr, column=1, value=f'=IF(Equipe!A{5 + i}="","",Equipe!A{5 + i})').font = f_link
    ws.cell(row=rr, column=2, value=f'=IF(A{rr}="","",SUMIFS({LT_MIN},{LT_ANAL},A{rr})/60)').number_format = NUM1
    ws.cell(row=rr, column=3, value=f'=IF(A{rr}="","",COUNTIF({CAD_RESP},A{rr}))').number_format = NUM0
    ws.cell(row=rr, column=4, value=(
        f'=IF(A{rr}="","",IF(C{rr}=0,"-",B{rr}*$B$6/C{rr}))')).number_format = NUM1
    for ci in range(1, 5):
        ws.cell(row=rr, column=ci).border = borda
AN = A0 + N_EQUIPE - 1

ws.cell(row=AN + 2, column=1, value="Tempo por classe de empresa").font = f_secao
cabecalho(ws, AN + 3, ["Classe", "Nº de empresas", "Horas no período", "Horas/empresa/mês", "", ""])
C0 = AN + 4
for i, cl in enumerate(["A", "B", "C"]):
    rr = C0 + i
    ws.cell(row=rr, column=1, value=cl).alignment = Alignment(horizontal="center")
    ws.cell(row=rr, column=2, value=f'=COUNTIF({CAD_CLASSE},A{rr})').number_format = NUM0
    ws.cell(row=rr, column=3, value=f'=SUMIFS({LT_MIN},{LT_CLASSE},A{rr})/60').number_format = NUM1
    ws.cell(row=rr, column=4, value=f'=IF(B{rr}=0,"-",C{rr}*$B$6/B{rr})').number_format = NUM1
    for ci in range(1, 5):
        ws.cell(row=rr, column=ci).border = borda
ws.cell(row=C0 + 4, column=1, value=(
    "Leitura: horas/empresa/mês por classe é o principal KPI do projeto. A meta de 12 meses é reduzi-lo em ~30%.")).font = f_sub
ws.freeze_panes = "A4"

# ======================================================================
# 6. Business Case
# ======================================================================
ws = wb.create_sheet("Business Case")
titulo(ws, "Business Case — automação do departamento fiscal",
       "Texto azul com fundo amarelo = premissa editável (estimativa inicial). Texto preto = fórmula. Texto verde = vínculo.")
larguras(ws, {"A": 62, "B": 18, "C": 18, "D": 14, "E": 14, "F": 14, "G": 16, "H": 18, "I": 12})


def premissa(cel, valor, fmt=None, nota=None):
    ws[cel] = valor
    ws[cel].font, ws[cel].fill, ws[cel].border = f_input, fill_key, borda
    if fmt:
        ws[cel].number_format = fmt
    if nota:
        ws[cel].comment = Comment(nota, "Plano de automação")


ws["A4"] = "Premissas da equipe"
ws["A4"].font = f_secao
ws["A5"], ws["A6"], ws["A7"], ws["A8"] = (
    "Analistas fiscais (nº)",
    "Custo mensal médio por analista: salário + encargos + benefícios (R$)",
    "Horas produtivas por analista por mês",
    "Cenário de automação",
)
premissa("B5", 12, NUM0, "Informado por você: 12 pessoas no departamento fiscal.")
premissa("B6", 5500, BRL, "Estimativa inicial. Substitua pela média real da folha do fiscal (a aba Equipe calcula ao lado).")
premissa("B7", 140, NUM0, "≈ 21 dias úteis × 8 h × 83% de tempo produtivo.")
premissa("B8", "Base", None, "Escolha: Conservador, Base ou Otimista.")
lista(ws, '"Conservador,Base,Otimista"', "B8")
ws["C6"] = f"={EQUIPE_MEDIA}"
ws["C6"].font, ws["C6"].number_format = f_link, BRL
ws["D6"] = "← média da aba Equipe (se preenchida)"
ws["D6"].font = f_sub

ws["A10"] = "Onde está o tempo e quanto pode ser automatizado em 12 meses"
ws["A10"].font = f_secao
hdr = ["Etapa", "% do tempo (premissa)", "% medido (Resumo Baseline)", "Conservador", "Base", "Otimista",
       "Potencial no cenário", "Horas liberadas/mês"]
cabecalho(ws, 11, hdr, altura=34)
B0 = 12
for i, (etapa, _) in enumerate(ETAPAS):
    rr = B0 + i
    ws.cell(row=rr, column=1, value=etapa).font = f_norm
    c = ws.cell(row=rr, column=2, value=PCT_TEMPO[i])
    c.font, c.fill, c.number_format = f_input, fill_key, PCT
    c = ws.cell(row=rr, column=3, value=f"='Resumo Baseline'!D{RESUMO_PCT_ROW0 + i}")
    c.font, c.number_format = f_link, PCT
    for j, cen in enumerate(["Conservador", "Base", "Otimista"]):
        c = ws.cell(row=rr, column=4 + j, value=POTENCIAL[cen][i])
        c.font, c.fill, c.number_format = f_input, fill_key, PCT
    ws.cell(row=rr, column=7, value=f"=INDEX(D{rr}:F{rr},1,MATCH($B$8,$D$11:$F$11,0))").number_format = PCT
    ws.cell(row=rr, column=8, value=f"=$B$5*$B$7*B{rr}*G{rr}").number_format = NUM1
    for ci in range(1, 9):
        ws.cell(row=rr, column=ci).border = borda
BT = B0 + len(ETAPAS)
ws.cell(row=BT, column=1, value="Total / ganho ponderado").font = f_bold
ws.cell(row=BT, column=2, value=f"=SUM(B{B0}:B{BT - 1})").number_format = PCT
ws.cell(row=BT, column=3, value=f"=SUM(C{B0}:C{BT - 1})").number_format = PCT
for j in range(3):
    col = get_column_letter(4 + j)
    ws.cell(row=BT, column=4 + j, value=f"=SUMPRODUCT($B${B0}:$B${BT - 1},{col}{B0}:{col}{BT - 1})").number_format = PCT
ws.cell(row=BT, column=7, value=f"=SUMPRODUCT(B{B0}:B{BT - 1},G{B0}:G{BT - 1})").number_format = PCT
ws.cell(row=BT, column=8, value=f"=SUM(H{B0}:H{BT - 1})").number_format = NUM1
for ci in range(1, 9):
    ws.cell(row=BT, column=ci).font = f_bold
    ws.cell(row=BT, column=ci).border = borda
ws.cell(row=BT + 1, column=1, value="Checagem da soma dos % do tempo").font = f_sub
ws.cell(row=BT + 1, column=2, value=f'=IF(ABS(B{BT}-1)<0.001,"OK","Ajuste: soma ≠ 100%")').font = f_bold
ws.cell(row=BT + 2, column=1, value=(
    "Quando o levantamento terminar, copie os % da coluna C para a coluna B.")).font = f_sub

R0 = BT + 4
ws.cell(row=R0, column=1, value="Custos e receitas do ecossistema").font = f_secao
itens_custo = [
    ("Analista de automação/dados: custo mensal (R$)", 7500,
     "Novo profissional ou alguém da equipe realocado (nesse caso, use o custo atual da pessoa)."),
    ("Ferramentas e licenças: captura de XML, workflow, BI (R$/mês)", 2500, "Cotar com 2 fornecedores na Fase 0."),
    ("Infraestrutura: servidor ou nuvem, banco, backup (R$/mês)", 800, "Estimativa."),
    ("APIs e IA: Integra Contador, Claude (R$/mês)", 800, "Cobrança por consumo; validar com volumes reais."),
    ("Nova receita com serviços habilitados: recuperação de monofásico, revisão de regime (R$/mês)", 0,
     "Opcional. Deixe 0 para análise conservadora."),
    ("Implantação: consultoria/desenvolvimento externo, valor único (R$)", 40000, "Estimativa para a fundação (Fases 1 e 2)."),
    ("Meses para diluir a implantação", 4, "Pagamento distribuído nos primeiros meses."),
]
for i, (rot, val, nota) in enumerate(itens_custo):
    rr = R0 + 1 + i
    ws.cell(row=rr, column=1, value=rot).font = f_norm
    premissa(f"B{rr}", val, NUM0 if "Meses" in rot else BRL, nota)
C_ANALISTA, C_FERR, C_INFRA, C_API, C_RECEITA, C_IMPL, C_MESES = (f"$B${R0 + 1 + i}" for i in range(7))

P0 = R0 + 1 + len(itens_custo) + 1
ws.cell(row=P0, column=1, value="Rampa de ganho (% do ganho pleno alcançado)").font = f_secao
rampa = [("Meses 1 a 3 (Fase 0 e início da Fase 1)", 0.10), ("Meses 4 a 6", 0.40),
         ("Meses 7 a 9", 0.70), ("Mês 10 em diante", 1.0)]
for i, (rot, val) in enumerate(rampa):
    rr = P0 + 1 + i
    ws.cell(row=rr, column=1, value=rot).font = f_norm
    premissa(f"B{rr}", val, PCT)
RAMPA_RNG = f"$B${P0 + 1}:$B${P0 + 4}"

S0 = P0 + 6
ws.cell(row=S0, column=1, value="Resultado").font = f_secao
FL0 = S0 + 11  # início do fluxo mensal (definido abaixo)
resultado = [
    ("Ganho de produtividade ponderado (cenário escolhido)", f"=G{BT}", PCT),
    ("Horas liberadas por mês em regime pleno", f"=H{BT}", NUM1),
    ("FTE equivalente liberado", f"=IF($B$7=0,0,B{S0 + 2}/$B$7)", '0.0 "FTE"'),
    ("Economia bruta anual se a capacidade virar não reposição de vagas (R$)", f"=B{S0 + 3}*$B$6*12", BRL),
    ("Nova receita anual (R$)", f"={C_RECEITA}*12", BRL),
    ("Custo recorrente anual do ecossistema (R$)", f"=({C_ANALISTA}+{C_FERR}+{C_INFRA}+{C_API})*12", BRL),
    ("Resultado líquido anual em regime pleno (R$)", f"=B{S0 + 4}+B{S0 + 5}-B{S0 + 6}", BRL),
    ("Payback: mês em que o acumulado fica positivo",
     f'=IF(COUNT(I{FL0 + 1}:I{FL0 + 24})=0,"acima de 24 meses",MIN(I{FL0 + 1}:I{FL0 + 24}))', '0 "º mês"'),
]
for i, (rot, form, fmt) in enumerate(resultado):
    rr = S0 + 1 + i
    ws.cell(row=rr, column=1, value=rot).font = f_bold if i in (2, 6, 7) else f_norm
    c = ws.cell(row=rr, column=2, value=form)
    c.number_format, c.border = fmt, borda
    c.font = f_bold if i in (2, 6, 7) else f_norm
ws.cell(row=S0 + 9, column=1, value=(
    "A economia só vira caixa se a capacidade liberada for convertida: não repor desligamentos, "
    "absorver novos clientes sem contratar ou vender novos serviços.")).font = f_sub

ws.cell(row=FL0 - 1, column=1, value="Fluxo mensal (24 meses)").font = f_secao
cabecalho(ws, FL0, ["Mês", "Rampa", "Economia (R$)", "Nova receita (R$)", "Custos recorrentes (R$)",
                    "Implantação (R$)", "Resultado do mês (R$)", "Acumulado (R$)", "Payback?"])
for m in range(1, 25):
    rr = FL0 + m
    ws.cell(row=rr, column=1, value=m if m == 1 else f"=A{rr - 1}+1")
    ws.cell(row=rr, column=2, value=f"=INDEX({RAMPA_RNG},MIN(4,ROUNDUP(A{rr}/3,0)))").number_format = PCT
    ws.cell(row=rr, column=3, value=f"=$B${S0 + 3}*$B$6*B{rr}").number_format = BRL
    ws.cell(row=rr, column=4, value=f"={C_RECEITA}*B{rr}").number_format = BRL
    ws.cell(row=rr, column=5, value=f"={C_ANALISTA}+{C_FERR}+{C_INFRA}+{C_API}").number_format = BRL
    ws.cell(row=rr, column=6, value=f"=IF(A{rr}<={C_MESES},{C_IMPL}/{C_MESES},0)").number_format = BRL
    ws.cell(row=rr, column=7, value=f"=C{rr}+D{rr}-E{rr}-F{rr}").number_format = BRL
    ws.cell(row=rr, column=8, value=(f"=G{rr}" if m == 1 else f"=H{rr - 1}+G{rr}")).number_format = BRL
    ws.cell(row=rr, column=9, value=f'=IF(H{rr}>=0,A{rr},"")').number_format = NUM0
    for ci in range(1, 10):
        ws.cell(row=rr, column=ci).border = borda
        if m % 2 == 0:
            ws.cell(row=rr, column=ci).fill = fill_band
ws.conditional_formatting.add(f"H{FL0 + 1}:H{FL0 + 24}", CellIsRule(
    operator="greaterThanOrEqual", formula=["0"], font=Font(name=FONTE, color="1D6B4F", bold=True)))
ws.freeze_panes = "B4"

# ======================================================================
# 7. Catálogo de Auditorias
# ======================================================================
ws = wb.create_sheet("Catálogo de Auditorias")
cat = json.load(open(CATALOGO_JSON, encoding="utf-8"))
titulo(ws, "Catálogo de Auditorias Fiscais (regras AUD)",
       "Prioridade = risco × 2 + (4 − esforço): de 3 a 9, quanto maior, antes. Status e responsável servem para acompanhar a construção.")
cols = [("Código", 10), ("Grupo", 24), ("Regra", 30), ("O que verifica", 60), ("Fontes cruzadas", 30),
        ("Regimes", 10), ("Tributos", 16), ("Risco", 8), ("Esforço", 9), ("Fase", 6), ("Frequência", 12),
        ("Ação", 12), ("Referência", 30), ("Prioridade", 10), ("Status", 16), ("Responsável", 20),
        ("Observações", 30)]
cabecalho(ws, 4, [c for c, _ in cols], altura=32)
for i, (_, w) in enumerate(cols, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w
for k, r in enumerate(cat["regras"]):
    rr = 5 + k
    vals = [r["cod"], r["grupo"], r["nome"], r["verifica"], r["fontes"], r["regimes"], r["tributos"],
            r["risco"], r["esforco"], r["fase"], r["freq"], r["acao"], r["ref"]]
    for ci, v in enumerate(vals, start=1):
        c = ws.cell(row=rr, column=ci, value=v)
        c.font, c.alignment, c.border = f_norm, wrap, borda
    ws.cell(row=rr, column=14, value=(
        f'=IF(H{rr}="Alto",3,IF(H{rr}="Médio",2,1))*2+4-IF(I{rr}="Alto",3,IF(I{rr}="Médio",2,1))'))
    ws.cell(row=rr, column=14).alignment = Alignment(horizontal="center", vertical="top")
    ws.cell(row=rr, column=10).alignment = Alignment(horizontal="center", vertical="top")
    for ci in (15, 16, 17):
        c = ws.cell(row=rr, column=ci)
        c.fill, c.border, c.font, c.alignment = fill_input, borda, f_norm, wrap
    ws.cell(row=rr, column=15, value="Backlog")
    ws.cell(row=rr, column=14).border = borda
    if k % 2 == 1:
        for ci in range(1, 15):
            ws.cell(row=rr, column=ci).fill = fill_band
UR = 4 + len(cat["regras"])
lista(ws, '"Backlog,Em especificação,Em desenvolvimento,Modo sombra,Em produção,Descartada"', f"O5:O{UR}")
lista(ws, f"={EQUIPE_RNG}", f"P5:P{UR}")
ws.conditional_formatting.add(f"H5:H{UR}", CellIsRule(operator="equal", formula=['"Alto"'],
                                                       font=Font(name=FONTE, color="9C0006", bold=True)))
ws.conditional_formatting.add(f"L5:L{UR}", CellIsRule(operator="equal", formula=['"Oportunidade"'],
                                                       font=Font(name=FONTE, color="1D6B4F", bold=True)))
ws.freeze_panes = "D5"
ws.auto_filter.ref = f"A4:Q{UR}"

# ======================================================================
# 8. Plano 30 dias
# ======================================================================
ws = wb.create_sheet("Plano 30 dias")
titulo(ws, "Plano da Fase 0 — 30 dias",
       "Informe a data de início (segunda-feira). Os prazos se ajustam sozinhos.")
ws["A4"] = "Data de início (segunda-feira)"
ws["A4"].font = f_bold
ws["C4"] = _dt.datetime(2026, 9, 28)
ws["C4"].number_format = "DD/MM/YYYY"
ws["C4"].font, ws["C4"].fill, ws["C4"].border = f_input, fill_key, borda
cabecalho(ws, 6, ["Semana", "Nº", "Ação", "Responsável (sugestão)", "Entregável", "Prazo", "Status",
                  "Observações"])
larguras(ws, {"A": 9, "B": 5, "C": 62, "D": 24, "E": 36, "F": 12, "G": 14, "H": 30})
plano = [
    (1, "URGENTE (até 30/09/2026): revisar clientes do Simples com vendas relevantes para empresas e decidir a opção pelo IBS/CBS fora do DAS no 1º semestre de 2027. A opção pode ser cancelada até 30/11/2026. Confirme as regras no Portal do Simples Nacional.", "Sócio + coordenador", "Clientes analisados e opções feitas"),
    (1, "Nomear o núcleo do projeto: patrocinador (sócio), dono das regras (coordenador fiscal) e candidato a analista de automação.", "Sócio", "Núcleo nomeado"),
    (1, "Comunicar a equipe: objetivos (qualidade, escala, menos retrabalho), como o tempo será medido e o compromisso com as pessoas.", "Sócio", "Reunião realizada"),
    (1, "Preencher a aba Equipe e iniciar o Levantamento de Tempo (10 dias úteis).", "Coordenador fiscal", "Registros diários"),
    (1, "Agendar reunião com o gerente de contas da Thomson Reuters (roteiro na aba Roteiro de Reuniões).", "Sócio", "Reunião agendada"),
    (2, "Preencher o Cadastro Mestre de 100% das empresas (cada analista, a sua carteira).", "Analistas", "Cadastro completo"),
    (2, "Inventariar certificados (A1/A3), procurações do e-CAC e acessos estaduais/municipais.", "Coordenador fiscal", "Pendências listadas"),
    (2, "Mapear o processo atual de 3 empresas-tipo: Simples comércio, Presumido serviços e regime normal de ICMS.", "Coordenador + 1 analista", "Fluxo atual desenhado"),
    (2, "Iniciar a campanha autXML: pedir aos clientes que incluam o CNPJ do escritório no campo autXML das NF-e que emitem.", "Coordenador + atendimento", "Comunicado enviado"),
    (2, "Reunião com a Thomson Reuters: registrar as respostas no roteiro.", "Sócio", "Roteiro respondido"),
    (3, "Pedir a chave da API da Domínio (api.dominio@tr.com), enviar XML de 1 empresa e testar a extração (EFD, Gerador de Relatórios e, se liberado, backup do banco).", "Analista de automação", "Envio e extração testados"),
    (3, "Contratar e testar o Integra Contador (SERPRO) com 2 empresas-piloto (procuração).", "Analista de automação", "Consulta PGDAS-D e caixa postal ok"),
    (3, "POCs: captura (2 parceiros homologados da API Domínio) e auditoria (Kolossus Auditor e 1 alternativa), usando o Catálogo de Auditorias como checklist de cobertura.", "Coordenador fiscal", "Comparativo preenchido"),
    (3, "Selecionar 15 a 20 empresas-piloto (mistura de classes A, B e C).", "Coordenador fiscal", "Lista do piloto"),
    (4, "Consolidar o baseline e atualizar o Business Case com os números reais.", "Sócio + coordenador", "Baseline aprovado"),
    (4, "Decidir: execução (interno × contratado), o que comprar × construir, infraestrutura (local × nuvem) e orçamento.", "Sócio", "Decisões registradas"),
    (4, "Preparar o ambiente: servidor/VM, banco de dados, cofre de certificados e acesso ao repositório.", "Analista de automação", "Ambiente pronto"),
    (4, "Aprovar o backlog da Fase 1 (16 regras) e o calendário de fechamento do piloto.", "Sócio + coordenador", "Kick-off da Fase 1"),
]
for i, (sem, acao, resp, entreg) in enumerate(plano):
    rr = 7 + i
    prazo = _dt.datetime(2026, 9, 30) if acao.startswith("URGENTE") else f"=$C$4+4+7*(A{rr}-1)"
    vals = [sem, i + 1, acao, resp, entreg, prazo, "A fazer", "Prazo legal fixo" if acao.startswith("URGENTE") else ""]
    for ci, v in enumerate(vals, start=1):
        c = ws.cell(row=rr, column=ci, value=v)
        c.font, c.alignment, c.border = f_norm, wrap, borda
    ws.cell(row=rr, column=6).number_format = "DD/MM/YYYY"
    for ci in (7, 8):
        ws.cell(row=rr, column=ci).fill = fill_input
    if sem % 2 == 0:
        for ci in range(1, 7):
            ws.cell(row=rr, column=ci).fill = fill_band
lista(ws, '"A fazer,Em andamento,Concluído,Bloqueado"', f"G7:G{6 + len(plano)}")
ws.freeze_panes = "A7"

# ======================================================================
# 9. Roteiro de Reuniões
# ======================================================================
ws = wb.create_sheet("Roteiro de Reuniões")
titulo(ws, "Roteiro de reuniões da Fase 0",
       "Perguntas que decidem a arquitetura. Registre a resposta e quem respondeu.")
cabecalho(ws, 4, ["Com quem", "Nº", "Pergunta", "Por que importa", "Resposta", "Respondido por / data"])
larguras(ws, {"A": 20, "B": 5, "C": 62, "D": 46, "E": 46, "F": 20})
ROTEIRO = [
    ("Thomson Reuters", "API Onvio BR Accounting (Integração com ERP): custo, limites de volume e como ativar a chave de integração em todas as empresas. As Rotinas Automáticas importam sem usuário logado na versão Web?",
     "É a porta oficial de entrada de XML na Escrita Fiscal, sem robô de tela."),
    ("Thomson Reuters", "Existe (ou está no roadmap) alguma API para LER dados: notas, itens, apurações, saldos?",
     "Hoje a API só envia dados para dentro da Domínio."),
    ("Thomson Reuters", "Podemos restaurar os backups do nosso banco (Suporte > Backup) em servidor próprio e ler via ODBC com usuário externo? Há exigência de licença do SQL Anywhere? Qual a frequência dos backups de modificações?",
     "É o caminho usado por fornecedores de BI para ter todos os dados da escrituração."),
    ("Thomson Reuters", "Leiaute Importação Padrão (TXT com separador): documentação atual; aceita acumulador e CFOP por nota? Funciona igual na versão Web?",
     "Permite lançar documentos sem XML ou com classificação definida pelo nosso motor."),
    ("Thomson Reuters", "Dá para gerar EFD e relatórios do Gerador de Relatórios em lote (várias empresas) e salvar em pasta mapeada?",
     "Plano B de extração se o backup não for liberado."),
    ("Thomson Reuters", "É permitido usar robôs (RPA) na Domínio Web? Exige usuário/licença dedicada? Como tratar o MFA do login?",
     "A Domínio Web é transmitida por imagem (GraphOn): robô de tela é frágil e precisa de regra clara."),
    ("Thomson Reuters", "Kolossus Auditor: quais cruzamentos cobre, custo por empresa, e se exporta os achados (Excel/API) para o nosso painel.",
     "Pode cobrir parte do catálogo sem desenvolvimento."),
    ("Thomson Reuters", "Domínio Processos e Portal do Cliente estão no nosso pacote (One/Pro/Max)? Como encerram tarefas automaticamente?",
     "Evita contratar outra ferramenta de workflow."),
    ("Thomson Reuters", "Cronograma da Domínio para CBS/IBS em 2026–2027: escrituração, apuração, Simples com IBS/CBS fora do DAS.",
     "A virada de janeiro/2027 muda regras e telas."),
    ("Thomson Reuters", "Como somos avisados de mudanças de tela ou de leiaute nas atualizações?",
     "Robôs e importações quebram com mudanças não avisadas."),
    ("SERPRO", "Integra Contador: confirmar os serviços do nosso caso (PGDAS-D, DEFIS, DCTFWeb/MIT, Caixa Postal, SITFIS, Pagamentos, Parcelamentos, Eventos de Atualização).",
     "Automatiza consultas e transmissões federais por API oficial."),
    ("SERPRO", "Tabela de preços vigente por faixa (consulta, emissão, declaração), contrato e prazo de ativação.", "Entra no Business Case."),
    ("SERPRO", "Requisitos: e-CNPJ do escritório, procuração eletrônica de cada cliente por código de serviço.",
     "Define a preparação com os clientes."),
    ("Fornecedores de captura", "É parceiro homologado da API Domínio? Envia os XML direto para a Escrita Fiscal?",
     "Captura e entrada na Domínio sem passo manual."),
    ("Fornecedores de captura", "Cobertura: NF-e, NFC-e, CT-e, NFS-e (ADN e municípios fora do ADN), eventos de cancelamento e manifestação; notas emitidas pelo cliente (autXML).",
     "A SEFAZ não devolve ao emitente as próprias notas; captura incompleta gera falso 'faltante'."),
    ("Fornecedores de captura", "Captura com A1 do cliente, sem certificado (autXML) ou com A3? Como respeitam o limite da SEFAZ (bloqueio de 1 hora por consumo indevido)?",
     "Muitos clientes têm A3, que não serve para robôs."),
    ("Fornecedores de captura", "Exportação/API para o nosso banco: formatos, frequência, histórico de 5 anos.",
     "O hub de dados precisa receber tudo sem digitação."),
    ("Fornecedores de captura", "Preço por empresa/volume, contrato, LGPD (acordo de tratamento de dados) e onde os dados ficam.",
     "Custo e conformidade."),
]
for i, (quem, perg, pq) in enumerate(ROTEIRO):
    rr = 5 + i
    vals = [quem, i + 1, perg, pq, "", ""]
    for ci, v in enumerate(vals, start=1):
        c = ws.cell(row=rr, column=ci, value=v)
        c.font, c.alignment, c.border = f_norm, wrap, borda
    for ci in (5, 6):
        ws.cell(row=rr, column=ci).fill = fill_input
ws.freeze_panes = "A5"

for s in wb.worksheets:
    s.sheet_properties.tabColor = VERDE if s.title in ("Leia-me", "Business Case") else "8FB8A4"
wb.save(SAIDA)
print("ok", SAIDA)
